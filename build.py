#!/usr/bin/env python3
"""
Construye Sistema_Gestion_MP2026.html a partir de:
  - template.html  (interfaz + lógica de la aplicación)
  - source.xlsx    (libro Excel original: datos maestros y registros)

Uso:  python3 build.py
Salida: Sistema_Gestion_MP2026.html (archivo autónomo, sin servidor)
"""
import openpyxl, json, datetime, os, sys

XLSX = "source.xlsx"
TPL  = "template.html"
OUT  = "Sistema_Gestion_MP2026.html"

def cv(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return v

def extract():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    def rows(name, cols):
        ws = wb[name]; out = []
        for r in range(2, ws.max_row + 1):
            rec, empty = {}, True
            for key, col in cols:
                v = cv(ws.cell(r, col).value)
                if v not in (None, ""): empty = False
                rec[key] = v if v is not None else ""
            if not empty: out.append(rec)
        return out

    data = {}
    data['equipos'] = rows('Equipos', [('inventario',3),('equipo',4),('servicio',5),('unidad',6),
        ('ubicacion',7),('procedencia',8),('marca',9),('modelo',10),('serie',11),('anio',12),('clasificacion',13)])

    cat = wb['Catálogos']
    def coll(col, start=4):
        out = []
        for r in range(start, cat.max_row + 1):
            v = cat.cell(r, col).value
            if v not in (None, "") and not str(v).startswith('('):
                out.append(str(v))
        return out
    data['catalogos'] = {
        'ingenieros': coll(1), 'estadoEquipo': coll(2), 'tipoCompra': coll(3),
        'resultadoVisita': coll(4), 'estadoEvento': coll(5), 'decisionReeval': coll(6),
        'periodicidad': coll(11), 'resultadoMP': coll(13), 'ejecucionMP': coll(16),
        'empresas': [cat.cell(r,1).value for r in range(17,22) if cat.cell(r,1).value],
        'servicios': coll(3, 17),
    }
    data['catalogos']['resultadoMPdesc'] = {str(cat.cell(r,13).value): cat.cell(r,14).value
        for r in range(4,18) if cat.cell(r,13).value}
    data['catalogos']['causales'] = {str(cat.cell(r,18).value): cat.cell(r,19).value
        for r in range(4,12) if cat.cell(r,18).value}

    data['ordenes'] = rows('1. Órdenes de Trabajo', [('folio',1),('fechaOT',2),('tecnico',3),
        ('requerimiento',4),('inventario',6),('estadoInicial',8),('estadoEvento',10),('fechaCierre',11)])
    data['compras'] = rows('2. Línea de Compra', [('folio',1),('ruta',2),('fechaCot',3),('proveedor',4),
        ('nCot',5),('monto',6),('tipoCompra',7),('nInforme',8),('fechaInforme',9),('respInforme',10),
        ('nOC',11),('fechaOC',12),('descOC',13),('fechaEnvioProv',14)])
    data['repuestos'] = rows('3. Repuestos', [('folio',1),('fechaRep',2),('descripcion',3),('estadoFinal',4)])
    data['visitas'] = rows('4. Visitas Técnicas', [('folio',1),('fechaVisita',2),('proveedor',3),('ingeniero',4),
        ('nReporte',5),('fechaReporte',6),('archivado',7),('reparo',8),('descripcion',9),('estadoFinal',10),
        ('fechaCierre',11),('reporteRecibido',12),('respConseguir',13),('respAdmin',14)])
    data['servicio'] = rows('5. Servicio Técnico (Envíos)', [('folio',1),('nHoja',2),('respEnvia',3),('fechaEnvio',4),
        ('empresa',5),('nReporte',6),('fechaReporte',7),('archivado',8),('fechaRetorno',9),('folioGuia',10),
        ('estadoRetorno',11),('decision',12),('reporteRecibido',13),('respConseguir',14),('respAdmin',15)])
    data['bajas'] = rows('6. Bajas', [('folio',1),('folioDoc',2),('respInterno',3),('fechaBaja',4),('motivo',5)])
    data['terreno'] = rows('7. Reparación en Terreno', [('folio',1),('fechaRep',2),('descripcion',3),('estadoFinal',4)])
    data['pendientes'] = rows('Pendientes', [('equipo',1),('serie',2),('inventario',3),('tipo',4),
        ('fechaCompromiso',5),('situacion',7),('estado',8),('prioridad',9),('respEjecutivo',10),
        ('respAdmin',11),('enEspera',12),('descripcion',13),('tareas',14),('fechaCompletado',15),('creado',16)])
    data['registroMP'] = rows('Registro MP', [('inventario',1),('periodo',7),('ejecucion',8),('empresa',9),
        ('resultado',10),('fechaEjec',12),('estadoEquipo',13),('ingeniero',14),('protInterno',15),
        ('protEmpresa',16),('observaciones',17)])

    ws = wb['MP Carta 2026']
    meses = ['Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']
    mp = []
    for r in range(4, ws.max_row + 1):
        inv = ws.cell(r, 1).value
        if inv in (None, ""): continue
        marks = {meses[i]: str(ws.cell(r, 9+i).value)
                 for i in range(12) if ws.cell(r, 9+i).value not in (None, "")}
        period = ws.cell(r, 8).value
        if marks or period:
            mp.append({'inventario': str(inv), 'periodicidad': period or '', 'marks': marks})
    data['mpCarta'] = mp
    return data

def main():
    if not os.path.exists(XLSX):
        sys.exit("Falta source.xlsx")
    data = extract()
    payload = json.dumps(data, ensure_ascii=False)
    assert '</script' not in payload.lower()
    tpl = open(TPL, encoding='utf-8').read()
    open(OUT, 'w', encoding='utf-8').write(tpl.replace('/*__SEED__*/{}', payload))
    print(f"OK -> {OUT} ({os.path.getsize(OUT):,} bytes)")
    for k, v in data.items():
        if isinstance(v, list): print(f"   {k}: {len(v)}")

if __name__ == "__main__":
    main()
